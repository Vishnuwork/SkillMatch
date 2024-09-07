from dotenv import load_dotenv

load_dotenv()
#its gonna load all the environment variables
import streamlit as st
from docx import Document
import os,io
import google.generativeai as genai
from PIL import Image
import pdf2image
import base64


genai.configure(api_key=os.getenv("GOOGLE_API_KEY"))

def get_docx(text):
    document = Document()
    document.add_paragraph(text)
    bio = io.BytesIO()
    document.save(bio)
    return bio.getvalue()

def get_genai_response(prompt, pdf_content, input_text):
        model=genai.GenerativeModel('gemini-1.5-flash')
        response = model.generate_content([prompt,pdf_content[0],input_text])
        if response and response.candidates and response.candidates[0].content.parts:
              return response.candidates[0].content.parts[0].text
        else:
            return "No content available in the response."
      

def input_pdf_setup(uploaded_file):
    if uploaded_file is not None:# convert the pdf to image
     images=pdf2image.convert_from_bytes(uploaded_file.read())

     first_page=images[0]

     #convert to bytes
     img_byte_arr=io.BytesIO()
     first_page.save(img_byte_arr,format="JPEG")
     img_byte_arr=img_byte_arr.getvalue()

     pdf_parts=[
        {
            "mime_type":"image/jpeg",
            "data": base64.b64encode(img_byte_arr).decode()#encode to base64
        }
    

     ]
     return pdf_parts
    else:
       raise FileNotFoundError("No File Uploaded")
     
## Streamlit App

st.set_page_config(page_title="Skill Match")
st.header("Skill Match")
input_text=st.text_area("Job Description: ",key="input")
uploaded_file=st.file_uploader("Upload Resume(PDF)",type=["pdf"])

if uploaded_file is not None:
   st.write("PDF Uploaded Successfully") 

submit1 = st.button("How's my Resume")
submit2 = st.button("Percentage Match")
submit3 = st.button("Technical Skills to ADD")
submit4 = st.button("Potential Projects")
submit5 = st.button("Beneficial Courses")
submit6 = st.button("Similar Companies")
submit7 = st.button("Cover Letter")

# Refined input prompts
input_prompt1 = """
You are an HR Technical Manager with expertise in Machine Learning, AI, Data Science, Software Development, DevOps, Cloud, Cybersecurity, Mobile App, Front-End, Back-End, and Blockchain. Evaluate the provided resume against the given job description, focusing on technical skills, experience, and qualifications. Highlight key strengths, identify weaknesses, and provide a professional assessment of how well the resume aligns with the job requirements.
"""

input_prompt2 = """
You are an ATS system designed to analyze resumes against job descriptions in fields like ML, AI, Data Science, DevOps, Cloud, Cybersecurity, Mobile App Development, Front-End, Back-End, and Blockchain. Evaluate the resume for a percentage match to the job description, considering skills, experience, education, and soft skills. Highlight both areas of alignment and areas needing improvement.
"""

input_prompt3 = """
As an HR Technical Manager specializing in ML, AI, Data Science, Software Development, DevOps, Cloud, Cybersecurity, Mobile App, Front-End, Back-End, and Blockchain, analyze the resume. Identify missing or weak technical skills relevant to the job description. Suggest specific programming languages, frameworks, cloud platforms, DevOps tools, or data tools that would strengthen the candidate’s profile.
"""

input_prompt4 = """
Based on the resume and job description, suggest 2-3 practical projects the candidate could undertake to demonstrate core technical skills. Focus on real-world problem-solving, teamwork, and innovation in areas like ML, AI, Software Development, DevOps, Cloud, Cybersecurity, Mobile App, Front-End, Back-End, and Blockchain. Projects should be directly aligned with the job requirements.
"""

input_prompt5 = """
Review the resume and identify gaps in technical skills relevant to the job description. Suggest specific courses or certifications to fill these gaps in areas like cloud platforms, DevOps tools, advanced programming, AI/ML, or cybersecurity. Recommend accredited platforms (Coursera, Udemy, etc.) and highlight relevant advanced certifications or courses.
"""

input_prompt6 = """
Based on the provided resume and job description, suggest multiple companies with similar roles. Focus on companies that use the same tech stack, operate in similar industries, or have a culture and mission aligned with the job requirements. Highlight companies where the candidate's skills would be a strong match.
"""

input_prompt7 = """
Generate a professional ATS-friendly cover letter based on the provided resume and job description. Incorporate the candidate's details directly without placeholders for job posting platforms or other extra spaces, addressing it to 'Dear Hiring Manager.' Ensure the formatting is optimized for ATS screening. Provide the output in a downloadable Word or PDF document format
"""
if submit1:
   if uploaded_file is not None:
      pdf_content=input_pdf_setup(uploaded_file)
      response=get_genai_response(input_prompt1,pdf_content,input_text)
      st.subheader(" You know my methods, Watson.")
      st.write(response)
   else:
      st.write("Please Upload the resume")


elif submit2:
   if uploaded_file is not None:
      pdf_content=input_pdf_setup(uploaded_file)
      response=get_genai_response(input_prompt2,pdf_content,input_text)
      st.subheader(" You know my methods, Watson.")
      st.write(response)
   else:
      st.write("Please Upload the resume")

elif submit3:
   if uploaded_file is not None:
      pdf_content=input_pdf_setup(uploaded_file)
      response=get_genai_response(input_prompt3,pdf_content,input_text)
      st.subheader(" You know my methods, Watson.")
      st.write(response)
   else:
      st.write("Please Upload the resume")

elif submit4:
   if uploaded_file is not None:
      pdf_content=input_pdf_setup(uploaded_file)
      response=get_genai_response(input_prompt4,pdf_content,input_text)
      st.subheader(" You know my methods, Watson.")
      st.write(response)
   else:
      st.write("Please Upload the resume")  

elif submit5:
   if uploaded_file is not None:
      pdf_content=input_pdf_setup(uploaded_file)
      response=get_genai_response(input_prompt5,pdf_content,input_text)
      st.subheader(" You know my methods, Watson.")
      st.write(response)
   else:
      st.write("Please Upload the resume")

elif submit6:
   if uploaded_file is not None:
      pdf_content=input_pdf_setup(uploaded_file)
      response=get_genai_response(input_prompt6,pdf_content,input_text)
      st.subheader(" You know my methods, Watson.")
      st.write(response)
   else:
      st.write("Please Upload the resume")

if submit7:
   if uploaded_file is not None:
      pdf_content=input_pdf_setup(uploaded_file)
      response=get_genai_response(input_prompt7,pdf_content,input_text)
      st.subheader(" You know my methods, Watson.")
      st.download_button("Download Cover Letter", data=get_docx(response), file_name="coverletter.docx", mime="docx")
   else:
      st.write("Please Upload the resume")










    